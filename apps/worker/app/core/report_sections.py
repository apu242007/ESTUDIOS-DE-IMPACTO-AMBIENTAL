"""Secciones de contenido del informe (después de datos generales): texto de catálogo, ambiente, impactos, PGA.

Todo sale de catálogos + lo que el profesional eligió/ajustó. Sin datos → "Sección incompleta" (A.6.9), nunca se inventa.
"""
from __future__ import annotations

import unicodedata
from typing import Any

from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.shared import Cm, Pt, RGBColor

from app.core.text_template import fill_vars, paragraphs, unresolved

STAGE_ORDER = ["construccion", "perforacion", "complementarias", "operacion", "abandono"]
STAGE_LABEL = {
    "construccion": "Construcción", "perforacion": "Perforación y terminación", "complementarias": "Obras complementarias",
    "operacion": "Operación", "abandono": "Abandono",
}
MEDIO_ORDER = ["fisico", "biotico", "perceptual", "cultural", "socioeconomico"]
MEDIO_LABEL = {
    "fisico": "Medio inerte", "biotico": "Medio biótico", "perceptual": "Medio perceptual",
    "cultural": "Medio sociocultural", "socioeconomico": "Medio socioeconómico",
}
AMBIENTE_ORDER = ["geologia", "suelos", "hidrologia", "clima", "flora", "fauna", "paisaje", "patrimonio", "socioeconomico", "otro"]


def norm(s: str) -> str:
    return unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower().strip()


def signed(n: float, d: int = 0) -> str:
    s = f"{abs(n):,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return ("+" if n > 0 else "−" if n < 0 else "") + s


# ----------------------------------------------------------------------------------------------- helpers de documento

def _pending(doc: Any, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"Sección incompleta: {text}")
    r.italic = True
    r.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)


def _write(doc: Any, text: str, vars: dict[str, Any], warn: list[str]) -> None:
    bad = unresolved(text, vars)
    if bad:
        warn.append("Variables sin resolver en un texto: " + ", ".join("{" + b + "}" for b in bad))
    for para in paragraphs(fill_vars(text, vars)):
        doc.add_paragraph(para)


def _table(doc: Any, header: list[str], rows: list[list[str]], widths: list[float], size: float = 8.5) -> Any:
    from app.jobs.docs import _repeat_header, _shade  # import tardío: docs.py importa este módulo

    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(size)
        _shade(c, "D9E2DF")
    _repeat_header(t.rows[0])
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(v).font.size = Pt(size)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


def chapter_blocks(ctx: dict[str, Any], chapter: str) -> list[tuple[str | None, str]]:
    """[(subtítulo o None, texto)] de los bloques 'seccion' de un capítulo, con el ajuste del proyecto si existe."""
    over = ctx.get("section_overrides") or {}
    out: list[tuple[str | None, str]] = []
    for b in sorted(ctx.get("sections") or [], key=lambda x: x.get("sort_order") or 0):
        title = b.get("title") or ""
        head, _, sub = title.partition(" / ")
        if norm(head) != norm(chapter):
            continue
        out.append((sub or None, over.get(b["key"], b["template"])))
    return out


def _chapter(doc: Any, ctx: dict[str, Any], chapter: str, warn: list[str], level: int = 2) -> bool:
    blocks = chapter_blocks(ctx, chapter)
    for sub, text in blocks:
        if sub:
            doc.add_heading(sub, level=level)
        _write(doc, text, ctx["vars"], warn)
    return bool(blocks)


# ----------------------------------------------------------------------------------------------- secciones

def add_resumen(doc: Any, ctx: dict[str, Any], n: int, warn: list[str]) -> None:
    doc.add_heading(f"{n}. Resumen ejecutivo", level=1)
    if not _chapter(doc, ctx, "RESUMEN EJECUTIVO", warn):
        _pending(doc, "no hay texto de resumen ejecutivo en el catálogo.")


def add_ubicacion(doc: Any, ctx: dict[str, Any], n: int, warn: list[str]) -> None:
    doc.add_heading(f"{n}. Ubicación y descripción general del proyecto", level=1)
    if not _chapter(doc, ctx, "UBICACION Y DESCRIPCION GENERAL DEL PROYECTO", warn):
        _pending(doc, "no hay textos de ubicación y descripción en el catálogo.")
    wells = ctx.get("wells") or []
    if wells:
        from app.core.geo import format_dms, to_gauss_kruger

        doc.add_heading("Coordenadas de los pozos", level=2)
        rows = []
        for w in wells:
            dlat, dlon = format_dms(w["lat"], w["lon"])
            x, y = to_gauss_kruger(w["lat"], w["lon"])
            rows.append([w["name"], dlat, dlon, str(round(x)), str(round(y))])
        _table(doc, ["Pozo", "Latitud", "Longitud", "X", "Y"], rows, [5, 3.2, 3.2, 2.3, 2.3])


def add_ambiente(doc: Any, ctx: dict[str, Any], n: int, warn: list[str]) -> None:
    doc.add_heading(f"{n}. Descripción general del ambiente", level=1)
    items = ctx.get("environment") or []
    if not ctx["project"].get("zone_key"):
        _pending(doc, "no se eligió la zona del proyecto.")
        return
    if not items:
        _pending(doc, "todos los apartados del ambiente fueron excluidos o no hay descripciones para la zona.")
        return
    for sec in AMBIENTE_ORDER:
        for it in [i for i in items if i["section"] == sec]:
            doc.add_heading(it["label"], level=2)
            _write(doc, it["body"], ctx["vars"], warn)
    for it in [i for i in items if i["section"] not in AMBIENTE_ORDER]:
        doc.add_heading(it["label"], level=2)
        _write(doc, it["body"], ctx["vars"], warn)


def _factor_stats(ctx: dict[str, Any]) -> dict[str, dict[str, Any]]:
    stats: dict[str, dict[str, Any]] = {}
    for i in ctx.get("impacts") or []:
        v = i.get("importance") or 0
        s = stats.setdefault(i["factor_id"], {"neg": 0, "pos": 0, "peor": None, "cat": None})
        if v < 0:
            s["neg"] += 1
            if s["peor"] is None or v < s["peor"]:
                s["peor"], s["cat"] = v, i.get("category")
        elif v > 0:
            s["pos"] += 1
    return stats


def add_impactos(doc: Any, ctx: dict[str, Any], n: int, warn: list[str]) -> None:
    doc.add_heading(f"{n}. Identificación de impactos y efectos ambientales", level=1)
    _chapter(doc, ctx, "IDENTIFICACION DE IMPACTOS", warn)
    impacts, factors = ctx.get("impacts") or [], ctx.get("factors") or []
    if not impacts or not factors:
        _pending(doc, "la matriz de impactos todavía no tiene celdas cargadas.")
        return
    doc.add_heading("Valoración de los impactos por factor", level=2)
    stats = _factor_stats(ctx)
    rows = []
    for medio in MEDIO_ORDER:
        for f in [x for x in factors if x["medio"] == medio]:
            s = stats.get(f["id"])
            rows.append([MEDIO_LABEL[medio], f["name"], str(s["neg"]) if s else "0", str(s["pos"]) if s else "0",
                         f"{signed(s['peor'])} {s['cat'] or ''}".strip() if s and s["peor"] is not None else "—"])
    _table(doc, ["Medio", "Factor", "Impactos negativos", "Impactos positivos", "Mayor impacto negativo"], rows, [3.4, 5.2, 2.3, 2.3, 3.2])
    doc.add_paragraph("La importancia se calcula como I = ±(3·IN + 2·EX + MO + PE + RV + SI + AC + EF + PR + MC). "
                      "La matriz completa figura en el anexo Análisis matricial.").runs[0].font.size = Pt(8.5)


def add_declaracion(doc: Any, ctx: dict[str, Any], n: int, warn: list[str]) -> None:
    doc.add_heading(f"{n}. Declaración de impacto ambiental", level=1)
    factors = ctx.get("factors") or []
    blocks = {b["key"]: b for b in (ctx.get("declarations") or [])}
    over = ctx.get("decl_overrides") or {}
    if not factors or not blocks:
        _pending(doc, "no hay declaraciones en el catálogo.")
        return
    faltan: list[str] = []
    hechos = 0
    for medio in MEDIO_ORDER:
        grupo = [f for f in factors if f["medio"] == medio]
        if not grupo:
            continue
        doc.add_heading(MEDIO_LABEL[medio], level=2)
        for f in grupo:
            b = blocks.get(f"decl_{f['code']}")
            text = over.get(f["id"]) or (b["template"] if b else None)
            if not text:
                faltan.append(f["name"])
                continue
            doc.add_heading((b or {}).get("title") or f["name"], level=3)
            _write(doc, text, ctx["vars"], warn)
            hechos += 1
    if faltan:
        _pending(doc, "faltan declaraciones para: " + ", ".join(faltan) + ".")
    if hechos == 0:
        warn.append("Ningún factor tiene declaración.")


def add_pga(doc: Any, ctx: dict[str, Any], n: int, warn: list[str]) -> None:
    doc.add_heading(f"{n}. Plan de gestión ambiental", level=1)
    generales, part = ctx.get("pga_general") or [], ctx.get("pga") or []
    if not generales and not part:
        _pending(doc, "no se eligieron medidas del PGA.")
        return
    if generales:
        doc.add_heading("Medidas generales", level=2)
        for t in generales:
            doc.add_paragraph(fill_vars(t, ctx["vars"]), style="List Bullet")
    if part:
        doc.add_heading("Medidas particulares", level=2)
        rows = []
        for stage in STAGE_ORDER + [None]:
            for m in [x for x in part if x.get("stage") == stage]:
                rows.append([STAGE_LABEL.get(stage or "", "—"), m.get("action") or "—", fill_vars(m["measure"], ctx["vars"]),
                             m.get("resource") or "—", m.get("timing") or "—", m.get("responsible") or "—", m.get("follow_up") or "—"])
        _table(doc, ["Etapa", "Acción", "Medida", "Recurso afectado", "Cronograma", "Responsable", "Seguimiento"], rows,
               [1.9, 2.2, 5.0, 2.2, 1.9, 1.9, 1.7], size=7.5)
    else:
        _pending(doc, "no se eligieron medidas particulares.")


def add_referencias(doc: Any, ctx: dict[str, Any], n: int, warn: list[str]) -> None:
    doc.add_heading(f"{n}. Referencias", level=1)
    if not _chapter(doc, ctx, "REFERENCIAS", warn):
        _pending(doc, "no hay referencias ni marco legal en el catálogo.")


# ----------------------------------------------------------------------------------------------- anexo: matriz (horizontal)

def _orientation(doc: Any, landscape: bool) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    sec.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = w, h
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(1.5 if landscape else 2.5))


def add_anexo_matriz(doc: Any, ctx: dict[str, Any], warn: list[str]) -> None:
    """Matriz completa factores × acciones en hoja horizontal. Vuelve a vertical al terminar."""
    factors, actions, impacts = ctx.get("factors") or [], ctx.get("actions") or [], ctx.get("impacts") or []
    _orientation(doc, True)
    doc.add_heading("Análisis matricial", level=2)
    if not impacts or not factors or not actions:
        _pending(doc, "la matriz de impactos todavía no tiene celdas cargadas.")
        _orientation(doc, False)
        return
    acts = sorted(actions, key=lambda a: (STAGE_ORDER.index(a["stage"]) if a.get("stage") in STAGE_ORDER else 99, a.get("sort_order") or 0))
    cells = {(i["factor_id"], i["action_id"]): i for i in impacts}
    header = ["Factor"] + [a["code"] for a in acts]
    rows = []
    for medio in MEDIO_ORDER:
        for f in [x for x in factors if x["medio"] == medio]:
            row = [f["name"]]
            for a in acts:
                c = cells.get((f["id"], a["id"]))
                row.append(signed(c["importance"]) if c and c.get("importance") is not None else "")
            rows.append(row)
    _table(doc, header, rows, [5.6] + [1.02] * len(acts), size=6.5)
    doc.add_paragraph()
    doc.add_paragraph("Acciones:").runs[0].bold = True
    for a in acts:
        p = doc.add_paragraph(f"{a['code']}  {a['name']}  ({STAGE_LABEL.get(a.get('stage') or '', 'sin etapa')})")
        p.runs[0].font.size = Pt(8)
    _orientation(doc, False)


# ----------------------------------------------------------------------------------------------- plantillas de cliente (docxtpl)

def template_data(ctx: dict[str, Any]) -> dict[str, Any]:
    """Marcadores extra para plantillas .docx de cliente: todo el contenido del informe como textos y listas.
    Ver docs/plantillas.md. Los textos ya vienen con las variables reemplazadas y con los ajustes del proyecto."""
    from app.core.geo import format_dms, to_gauss_kruger

    vars_ = ctx["vars"]
    med = lambda m: MEDIO_LABEL.get(m, m)  # noqa: E731

    def cap(name: str) -> list[dict[str, str]]:
        return [{"titulo": sub or "", "texto": fill_vars(t, vars_)} for sub, t in chapter_blocks(ctx, name)]

    stats = _factor_stats(ctx)
    factors = ctx.get("factors") or []
    actions = sorted(ctx.get("actions") or [], key=lambda a: (STAGE_ORDER.index(a["stage"]) if a.get("stage") in STAGE_ORDER else 99, a.get("sort_order") or 0))
    cells = {(i["factor_id"], i["action_id"]): i for i in ctx.get("impacts") or []}
    blocks = {b["key"]: b for b in ctx.get("declarations") or []}
    over = ctx.get("decl_overrides") or {}

    declaraciones = []
    for medio in MEDIO_ORDER:
        for f in [x for x in factors if x["medio"] == medio]:
            b = blocks.get(f"decl_{f['code']}")
            text = over.get(f["id"]) or (b["template"] if b else "")
            if text:
                declaraciones.append({"titulo": (b or {}).get("title") or f["name"], "medio": med(medio), "texto": fill_vars(text, vars_)})

    pozos = []
    for w in ctx.get("wells") or []:
        dlat, dlon = format_dms(w["lat"], w["lon"])
        x, y = to_gauss_kruger(w["lat"], w["lon"])
        pozos.append({"nombre": w["name"], "lat": dlat, "lon": dlon, "x": round(x), "y": round(y)})

    return {
        "borrador": not ctx.get("final"),
        "variables": vars_,
        "secciones": {"resumen": cap("RESUMEN EJECUTIVO"), "ubicacion": cap("UBICACION Y DESCRIPCION GENERAL DEL PROYECTO"),
                      "impactos": cap("IDENTIFICACION DE IMPACTOS"), "referencias": cap("REFERENCIAS")},
        "ambiente": [{"seccion": i["section"], "titulo": i["label"], "texto": fill_vars(i["body"], vars_)} for i in ctx.get("environment") or []],
        "factores": [{"codigo": f["code"], "nombre": f["name"], "medio": med(f["medio"]), "uip": f.get("uip"),
                      "componente": f.get("component") or "", "negativos": (stats.get(f["id"]) or {}).get("neg", 0),
                      "positivos": (stats.get(f["id"]) or {}).get("pos", 0), "peor": (stats.get(f["id"]) or {}).get("peor"),
                      "categoria_peor": (stats.get(f["id"]) or {}).get("cat") or ""} for f in factors],
        "matriz": {"acciones": [{"codigo": a["code"], "nombre": a["name"], "etapa": STAGE_LABEL.get(a.get("stage") or "", "")} for a in actions],
                   "filas": [{"factor": f["name"], "medio": med(f["medio"]),
                              "valores": [signed(cells[(f["id"], a["id"])]["importance"]) if (f["id"], a["id"]) in cells and cells[(f["id"], a["id"])].get("importance") is not None else "" for a in actions]}
                             for f in factors]},
        "declaraciones": declaraciones,
        "pga": {"generales": [fill_vars(t, vars_) for t in ctx.get("pga_general") or []],
                "particulares": [{"etapa": STAGE_LABEL.get(m.get("stage") or "", ""), "accion": m.get("action") or "",
                                  "medida": fill_vars(m["measure"], vars_), "recurso": m.get("resource") or "", "cronograma": m.get("timing") or "",
                                  "responsable": m.get("responsible") or "", "seguimiento": m.get("follow_up") or ""}
                                 for m in ctx.get("pga") or []]},
        "pozos": pozos,
    }
