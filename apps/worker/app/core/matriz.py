"""Lee la matriz de impactos (Excel tipo Conesa ponderada) y arma los datos para los catálogos. Sin red ni Supabase.

No inventa nada: lo que el Excel no trae (nombres de las opciones de cada atributo) queda vacío y va al reporte.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

ATTRS = ("IN", "EX", "MO", "PE", "RV", "SI", "AC", "EF", "PR", "MC")
ATTR_COLS = dict(zip(ATTRS, "DEFGHIJKLM"))  # columnas D..M de cada hoja de factor
ATTR_NAMES = {  # encabezados de la fila 1 de cada hoja de factor
    "IN": "Intensidad", "EX": "Extensión", "MO": "Momento", "PE": "Persistencia", "RV": "Reversibilidad",
    "SI": "Sinergia", "AC": "Acumulación", "EF": "Efecto", "PR": "Periodicidad", "MC": "Recuperabilidad",
}
# etapa del Excel → etapa del esquema (works.stage / catalog_impact_actions.stage)
STAGES = {
    "CONSTRUCCION": "construccion",
    "PERFORACION Y TERMINACION": "perforacion",
    "OBRAS COMPLEMENTARIAS": "complementarias",
    "OPERACION": "operacion",
    "ABANDONO": "abandono",
}
# medio del Excel (col B de la fila 2 de cada hoja + col A de "Matriz Ponderada") → catalog_impact_factors.medio
MEDIOS = {"INERTE": "fisico", "BIOTICO": "biotico", "PERCEPTUAL": "perceptual", "SOCIOCULTURAL": "cultural",
          "SOCIOECONOMICO": "socioeconomico"}


def norm(s: Any) -> str:
    """Mayúsculas sin tildes ni signos, para comparar rótulos del Excel."""
    t = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9 ]+", " ", t).strip()


def importance(sign: int, a: dict[str, int]) -> int:
    """I = signo × (3·IN + 2·EX + MO + PE + RV + SI + AC + EF + PR + MC). Igual que el trigger calc_impact."""
    return sign * (3 * a["IN"] + 2 * a["EX"] + a["MO"] + a["PE"] + a["RV"] + a["SI"] + a["AC"] + a["EF"] + a["PR"] + a["MC"])


@dataclass
class Cell:
    stage: str
    action: str
    sign: int
    attrs: dict[str, int]
    excel_i: int  # importancia que trae el Excel (para verificar la fórmula)


@dataclass
class FactorData:
    sheet: str
    name: str
    component: str
    medio: str
    uip: float | None
    scales: dict[str, str]  # texto de la escala de cada atributo tal como está en el Excel ("1-2-4-8-12")
    cells: list[Cell] = field(default_factory=list)


@dataclass
class Matriz:
    factors: list[FactorData]
    actions: list[tuple[str, str]]  # (etapa, nombre) en orden de aparición
    categories: list[dict[str, Any]]
    report: list[str]


def _num(v: Any) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _uip_by_factor(ws: Any) -> dict[str, float]:
    """UIP (unidades de importancia, suman 1000) de cada factor: columna E de "Matriz Ponderada"."""
    out: dict[str, float] = {}
    for r in range(5, ws.max_row + 1):
        name, uip = ws.cell(r, 4).value, _num(ws.cell(r, 5).value)
        if isinstance(name, str) and name.strip() not in ("", "0") and uip is not None and uip > 0:
            out[norm(name)] = uip
    return out


def parse_matriz(path: Path) -> Matriz:
    wb = openpyxl.load_workbook(path, data_only=True)
    report: list[str] = []
    uip = _uip_by_factor(wb["Matriz Ponderada"])
    factors: list[FactorData] = []
    actions: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()

    for ws in wb.worksheets:
        # hoja de factor = encabezado "SUBSISTEMA ... Signo" (las hojas resumen Acciones/Factores/Efectos no lo tienen)
        if str(ws["A1"].value).strip().upper() != "SUBSISTEMA" or str(ws["C1"].value).strip().lower() != "signo":
            continue
        if ws.sheet_state != "visible":
            report.append(f"Hoja oculta omitida: {ws.title!r} (no forma parte de la matriz publicada).")
            continue
        name = str(ws["B4"].value or "").strip()
        medio = MEDIOS.get(norm(ws["B2"].value))
        if not name or medio is None:
            report.append(f"Hoja {ws.title!r}: no se pudo leer el factor o el medio ({ws['B2'].value!r}); omitida.")
            continue
        fd = FactorData(ws.title, name, str(ws["B3"].value or "").strip(), medio, uip.get(norm(name)),
                        {a: str(ws[f"{col}3"].value or "") for a, col in ATTR_COLS.items()})
        if fd.uip is None:
            report.append(f"Factor {name!r}: sin UIP en 'Matriz Ponderada' (queda vacío).")
        stage = ""
        for r in range(5, ws.max_row + 1):
            a_val, action = ws.cell(r, 1).value, ws.cell(r, 2).value
            if isinstance(a_val, str) and norm(a_val) in STAGES:
                stage = STAGES[norm(a_val)]
            if not isinstance(action, str) or action.strip() in ("", "0") or not stage:
                continue
            sign = _num(ws.cell(r, 3).value)
            vals = {a: _num(ws[f"{col}{r}"].value) for a, col in ATTR_COLS.items()}
            action = " ".join(action.split())
            if (stage, action) not in seen:
                seen.add((stage, action))
                actions.append((stage, action))
            if not sign or any(v is None for v in vals.values()):
                continue  # celda que no aplica (signo 0): no es un impacto
            fd.cells.append(Cell(stage, action, int(sign), {a: int(v or 0) for a, v in vals.items()},
                                 int(_num(ws.cell(r, 14).value) or 0)))
        factors.append(fd)

    # Rangos de categoría: el Excel los tiene SOLO en el formato condicional de las hojas de factor
    # (−25/−50 moderado, −50/−75 severo, < −75 crítico). Bajo = lo que queda por debajo de 25 (por exclusión).
    categories = [
        {"applies_to": "negativo", "label": "Bajo", "min_abs": 0, "max_abs": 25},
        {"applies_to": "negativo", "label": "Moderado", "min_abs": 25, "max_abs": 50},
        {"applies_to": "negativo", "label": "Severo", "min_abs": 50, "max_abs": 75},
        {"applies_to": "negativo", "label": "Crítico", "min_abs": 75, "max_abs": None},
        {"applies_to": "positivo", "label": "Positivo", "min_abs": 0, "max_abs": None},
    ]
    report.append("Categorías: el Excel no define si |I| = 25, 50 o 75 cae en la categoría de arriba o de abajo "
                  "(el formato condicional usa rangos cerrados que se pisan). Se usó [25,50), [50,75), ≥75.")
    report.append("Opciones de atributos: el Excel trae solo los números de cada escala (fila 3), sin nombres. "
                  "Las etiquetas quedan iguales al número; el admin puede nombrarlas en Administración → Catálogos.")
    return Matriz(factors, actions, categories, report)


def attr_options(m: Matriz) -> dict[str, list[int]]:
    """Valores válidos de cada atributo: los de la escala del Excel + los que realmente se usan en las celdas.

    "(+4)" en las escalas de EX y MO es ambiguo (¿suma 4 o es el valor 12/8?): no se interpreta; si esos valores
    se usan de verdad en alguna celda entran igual por el segundo origen.
    """
    out: dict[str, set[int]] = {a: set() for a in ATTRS}
    for f in m.factors:
        for a in ATTRS:
            out[a].update(int(x) for x in re.findall(r"(?<![+\d])\d+(?!\d)", re.sub(r"\(\+?\d+\)", "", f.scales[a])))
            out[a].update(c.attrs[a] for c in f.cells)
    return {a: sorted(v) for a, v in out.items()}
