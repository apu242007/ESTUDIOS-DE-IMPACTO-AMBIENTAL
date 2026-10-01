"""Generación del informe (DOCX + PDF): contexto desde la base, fotos comprimidas, DOCX por plantilla o armado base."""
from __future__ import annotations

import datetime as dt
import io
import re
import shutil
import subprocess
import zipfile
import tempfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageOps

from app.core.geo import format_dms, to_gauss_kruger

BUCKET = "project-files"
DEFAULT_FIGURA = "Punto de interés"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre"]
TITULO = {
    "IA": "INFORME AMBIENTAL",
    "MTD": "MEMORIA TÉCNICA DESCRIPTIVA",
    "IA+MTD": "INFORME AMBIENTAL Y MEMORIA TÉCNICA DESCRIPTIVA",
}
KIND_LABEL = {
    "locacion": "Locación", "camino": "Camino", "ducto": "Ducto", "acueducto_temporal": "Acueducto temporal",
    "linea_electrica": "Línea eléctrica", "fibra_optica": "Fibra óptica", "predio": "Predio", "apendice": "Apéndice",
    "instalacion_aux": "Instalación auxiliar", "pozo_area": "Pozo / área de pozo",
}
Runner = Callable[..., Any]


# ----------------------------------------------------------------------------------------------- parámetros

def clean_params(raw: dict[str, Any] | None) -> dict[str, int]:
    """Calidad de las fotos del anexo. Se acota para que un valor absurdo no rompa ni infle el informe."""
    raw = raw or {}

    def num(key: str, default: int, lo: int, hi: int) -> int:
        try:
            v = int(raw.get(key, default))
        except (TypeError, ValueError):
            v = default
        return max(lo, min(hi, v))

    return {"photo_max_px": num("photo_max_px", 1200, 400, 3000), "jpeg_quality": num("jpeg_quality", 75, 40, 95)}


# ----------------------------------------------------------------------------------------------- interferencias

def render_template(tpl: str, vars: dict[str, Any]) -> str:
    """Plantilla con {variables}; sin IA (A.6.11). Mismo comportamiento que la web."""
    out = re.sub(r"\{(\w+)\}", lambda m: str(vars.get(m.group(1)) if vars.get(m.group(1)) is not None else ""), tpl)
    out = re.sub(r"\s+([,.;:])", r"\1", out)
    return re.sub(r"\s{2,}", " ", out).strip()


def build_interferencias(waypoints: list[dict[str, Any]], codes: dict[str, str], fichas: dict[str, int | None],
                         template: str | None = None) -> tuple[list[dict[str, Any]], int]:
    """Filas (Figura, Lat, Lon, X, Y, Cota, Descripción) de los waypoints con posición, y cuántos quedaron sin ella."""
    rows: list[dict[str, Any]] = []
    sin = 0
    for w in waypoints:
        lat, lon = w.get("lat"), w.get("lon")
        if lat is None or lon is None:
            sin += 1
            continue
        figura = codes.get(w.get("code") or "") or DEFAULT_FIGURA
        x, y = to_gauss_kruger(lat, lon)
        dlat, dlon = format_dms(lat, lon)
        obs = (w.get("description") or "").strip()
        if template:
            desc = render_template(template, {"figura": figura, "sigla": w.get("code"), "numero": w.get("number"),
                                              "vistas": w.get("views"), "observaciones": w.get("description")})
        else:
            desc = f"{figura}: {obs}" if obs else figura
            if w.get("views"):
                desc += f" (vistas: {w['views']})"
        ele = w.get("elevation_m")
        rows.append({"figura": figura, "lat": dlat, "lon": dlon, "x": round(x), "y": round(y),
                     "cota": None if ele is None else round(ele), "descripcion": desc,
                     "_lat": lat, "_lon": lon, "_ele": ele, "_ficha": fichas.get(w.get("line_id")) or 0, "_num": w.get("number") or 0})
    rows.sort(key=lambda r: (r["_ficha"], r["_num"]))
    return rows, sin


# ----------------------------------------------------------------------------------------------- fotos

def prepare_photo(data: bytes, max_px: int, quality: int) -> bytes:
    """Reduce y comprime para el anexo: respeta la orientación EXIF, no agranda y no arrastra metadatos."""
    with Image.open(io.BytesIO(data)) as im:
        im = ImageOps.exif_transpose(im).convert("RGB")
        im.thumbnail((max_px, max_px), Image.Resampling.LANCZOS)
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=quality, optimize=True)
        return buf.getvalue()


# ----------------------------------------------------------------------------------------------- DOCX base

def _shade(cell: Any, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _repeat_header(row: Any) -> None:
    pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    pr.append(el)


def _table(doc: Any, header: list[str], rows: list[list[str]], widths_cm: list[float] | None = None) -> Any:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        _shade(c, "D9E2DF")
    _repeat_header(t.rows[0])
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(v)
            run.font.size = Pt(9)
    if widths_cm:
        for row in t.rows:
            for i, w in enumerate(widths_cm):
                row.cells[i].width = Cm(w)
    return t


def _pending(doc: Any, text: str) -> None:
    """Sección sin datos: se marca 'incompleta' y se sigue (A.6.9), sin inventar contenido."""
    p = doc.add_paragraph()
    run = p.add_run(f"Sección incompleta: {text}")
    run.italic = True
    run.font.color.rgb = RGBColor(0x8A, 0x5A, 0x00)


def _page_number_footer(section: Any) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _fmt_m(v: Any, unit: str) -> str:
    return "—" if v is None else f"{float(v):,.1f} {unit}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fecha(today: dt.date) -> str:
    return f"{MESES[today.month - 1].capitalize()} de {today.year}"


def _draft_header(section: Any) -> None:
    """Encabezado condicional: las versiones no aprobadas salen marcadas (la final, no)."""
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BORRADOR — versión sin aprobar")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0xB4, 0x23, 0x18)


def _add_alcance(doc: Any, ctx: dict[str, Any]) -> None:
    doc.add_heading("Alcance de obras", level=2)
    if ctx["works"]:
        _table(doc, ["Obra", "Tipo", "Declarado", "Medido en el relevamiento"], [
            [w["name"], KIND_LABEL.get(w["kind"], w["kind"]),
             _fmt_m(w.get("declared_length_m"), "m") if w.get("declared_length_m") is not None
             else _fmt_m(w.get("declared_area_m2"), "m²"),
             _fmt_m(w.get("geom_length_m"), "m") if w.get("geom_length_m") is not None
             else _fmt_m(w.get("geom_area_m2"), "m²")] for w in ctx["works"]], [6, 3.5, 3, 3.5])
        sin_geom = sum(1 for w in ctx["works"] if w.get("geom_length_m") is None and w.get("geom_area_m2") is None)
        if sin_geom:
            _pending(doc, f"{sin_geom} obra(s) todavía sin geometría medida.")
    else:
        _pending(doc, "no se cargaron las obras del alcance.")


def _add_interferencias(doc: Any, ctx: dict[str, Any]) -> None:
    doc.add_heading("Interferencias y puntos de interés", level=2)
    if ctx["interferencias"]:
        _table(doc, ["Figura", "Latitud", "Longitud", "X", "Y", "Cota", "Descripción"],
               [[r["figura"], r["lat"], r["lon"], str(r["x"]), str(r["y"]), "" if r["cota"] is None else str(r["cota"]),
                 r["descripcion"]] for r in ctx["interferencias"]], [2.4, 2.4, 2.4, 1.7, 1.7, 1.1, 4.3])
        doc.add_paragraph("Coordenadas planas POSGAR 94 / Argentina faja 2 (EPSG:22182): X = norte, Y = este.").runs[0].font.size = Pt(8.5)
    else:
        _pending(doc, "no hay waypoints con posición.")
    from app.core.report_sections import add_figure

    add_figure(doc, ctx, "interferencias")
    if ctx["interferencias_sin_posicion"]:
        _pending(doc, f"{ctx['interferencias_sin_posicion']} waypoint(s) sin posición no figuran en la tabla.")


def _add_fotos(doc: Any, photos: list[dict[str, Any]]) -> None:
    doc.add_heading("Relevamiento fotográfico", level=2)
    if not photos:
        _pending(doc, "no hay fotos cargadas.")
    order: list[str] = []
    for ph in photos:
        if ph["category"] not in order:
            order.append(ph["category"])
    n_foto = 0
    for cat in order:
        group = [ph for ph in photos if ph["category"] == cat]
        doc.add_heading(group[0]["label"], level=3)
        for i in range(0, len(group), 2):
            tbl = doc.add_table(rows=1, cols=2)
            for j, ph in enumerate(group[i:i + 2]):
                n_foto += 1
                cell = tbl.rows[0].cells[j]
                cell.text = ""
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell.paragraphs[0].add_run().add_picture(io.BytesIO(ph["jpeg"]), width=Cm(7.6))
                cap = cell.add_paragraph(f"Foto {n_foto}." + (f" {ph['caption']}" if ph.get("caption") else ""))
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].font.size = Pt(8.5)


def _add_georef(doc: Any, ctx: dict[str, Any]) -> None:
    doc.add_heading("Archivos georreferenciados", level=2)
    files = ctx["layers"] + ctx["gps"]
    if files:
        _table(doc, ["Archivo", "Tipo", "Elementos / puntos"], [[f["name"], f["kind"], str(f.get("n") or "—")] for f in files], [8, 3, 5])
    else:
        _pending(doc, "no hay capas ni archivos de GPS importados.")


def build_docx(ctx: dict[str, Any], photos: list[dict[str, Any]], today: dt.date | None = None) -> bytes:
    """Informe completo (sin plantilla del cliente), en el orden del IA real. `photos` ya comprimidas.

    ctx["final"] = True → sin encabezado BORRADOR. Los avisos (variables sin resolver) quedan en ctx["_warn"].
    """
    from app.core import report_sections as rs
    from app.core.text_template import project_vars

    today = today or dt.date.today()
    p = ctx["project"]
    ctx.setdefault("vars", project_vars(p, (ctx.get("client") or {}).get("name")))
    warn: list[str] = ctx.setdefault("_warn", [])
    ctx["_fig_n"] = 0   # numeración de figuras propia de cada armado (no se arrastra entre llamadas)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.orientation = WD_ORIENT.PORTRAIT
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Cm(2.5))
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    _page_number_footer(sec)
    if not ctx.get("final"):
        _draft_header(sec)

    # --- carátula
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(TITULO.get(p["doc_type"], "INFORME AMBIENTAL"))
    r.bold = True
    r.font.size = Pt(22)
    n = doc.add_paragraph()
    n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = n.add_run(p["name"])
    r.bold = True
    r.font.size = Pt(16)
    for line in filter(None, [
        f"Código: {p['code']}" if p.get("code") else None,
        f"Yacimiento: {p['field_area']}" if p.get("field_area") else None,
        p.get("province"),
        f"Cliente: {ctx['client']['name']}" if ctx["client"].get("name") else None,
    ]):
        q = doc.add_paragraph(line)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(6):
        doc.add_paragraph()
    cons = p.get("consultant") or {}
    for line in filter(None, [cons.get("razon_social"), cons.get("responsable"), _fecha(today)]):
        q = doc.add_paragraph(line)
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- contenido (índice estático: sin números de página, para que sea igual en Word y en el PDF)
    capitulos = ["Datos generales", "Resumen ejecutivo", "Ubicación y descripción general del proyecto",
                 "Descripción general del ambiente", "Identificación de impactos y efectos ambientales",
                 "Declaración de impacto ambiental", "Plan de gestión ambiental", "Referencias", "Anexos"]
    doc.add_heading("Contenido", level=1)
    for i, c in enumerate(capitulos, 1):
        doc.add_paragraph(f"{i}. {c}")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- 1. datos generales
    doc.add_heading("1. Datos generales", level=1)
    app = p.get("applicant") or {}
    faltan = [n for n, v in (("razón social del solicitante", app.get("razon_social")),
                             ("consultora", cons.get("razon_social")),
                             ("responsable técnico", cons.get("responsable"))) if not v]
    _table(doc, ["Dato", "Detalle"], [
        ["Proyecto", p["name"]], ["Código", p.get("code") or "—"], ["Yacimiento / área", p.get("field_area") or "—"],
        ["Provincia", p.get("province") or "—"], ["Tipo de documento", p["doc_type"]],
        ["Empresa solicitante", app.get("razon_social") or "—"], ["CUIT", app.get("cuit") or "—"],
        ["Domicilio", app.get("domicilio") or "—"], ["Consultora", cons.get("razon_social") or "—"],
        ["Responsable técnico", cons.get("responsable") or "—"], ["Matrícula", cons.get("matricula") or "—"],
    ], [5, 11])
    if faltan:
        _pending(doc, "faltan " + ", ".join(faltan) + ".")

    # --- 2 a 8
    rs.add_resumen(doc, ctx, 2, warn)
    rs.add_ubicacion(doc, ctx, 3, warn)
    _add_alcance(doc, ctx)
    _add_interferencias(doc, ctx)
    rs.add_ambiente(doc, ctx, 4, warn)
    rs.add_impactos(doc, ctx, 5, warn)
    rs.add_declaracion(doc, ctx, 6, warn)
    rs.add_pga(doc, ctx, 7, warn)
    rs.add_referencias(doc, ctx, 8, warn)

    # --- 9. anexos
    doc.add_page_break()
    doc.add_heading("9. Anexos", level=1)
    _add_fotos(doc, photos)
    rs.add_anexo_matriz(doc, ctx, warn)   # hoja horizontal; vuelve a vertical
    _add_georef(doc, ctx)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_with_template(template: bytes, ctx: dict[str, Any], photos: list[dict[str, Any]]) -> bytes:
    """Plantilla .docx del cliente con variables docxtpl. Marcadores: ver docs/plantillas.md."""
    from docxtpl import DocxTemplate, InlineImage

    tpl = DocxTemplate(io.BytesIO(template))
    p = ctx["project"]
    data = {
        "proyecto": p, "cliente": ctx["client"], "solicitante": p.get("applicant") or {},
        "consultora": p.get("consultant") or {}, "obras": ctx["works"], "interferencias": ctx["interferencias"],
        "fotos": [{"categoria": ph["label"], "epigrafe": ph.get("caption") or "",
                   "imagen": InlineImage(tpl, io.BytesIO(ph["jpeg"]), width=Cm(8))} for ph in photos],
        "capas": ctx["layers"], "gps": ctx["gps"], "fecha": _fecha(dt.date.today()),
        "titulo": TITULO.get(p["doc_type"], "INFORME AMBIENTAL"),
    }
    from app.core import report_sections as rs
    from app.core.text_template import project_vars

    ctx.setdefault("vars", project_vars(p, (ctx.get("client") or {}).get("name")))
    data.update(rs.template_data(ctx))
    # plantilla subida por un usuario: Jinja en sandbox (sin acceso a __globals__/os) y con escape de &, <, >
    from jinja2.sandbox import SandboxedEnvironment

    tpl.render(data, jinja_env=SandboxedEnvironment(), autoescape=True)
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


# ----------------------------------------------------------------------------------------------- PDF

def find_soffice() -> str | None:
    for c in ("soffice", "libreoffice"):
        if (p := shutil.which(c)):
            return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe", r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(p).exists():
            return p
    return None


def docx_to_pdf(docx: Path, run: Runner = subprocess.run, exe: str | None = None) -> Path:
    """`soffice --headless --convert-to pdf`. Lanza RuntimeError con un mensaje legible si no se puede."""
    exe = exe or find_soffice()
    if not exe:
        raise RuntimeError("LibreOffice no está instalado en la PC del worker: se entrega solo el DOCX.")
    try:
        r = run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                capture_output=True, text=True, timeout=180)
    except subprocess.TimeoutExpired as e:
        raise RuntimeError("LibreOffice tardó demasiado en convertir a PDF.") from e
    pdf = docx.with_suffix(".pdf")
    if r.returncode != 0 or not pdf.exists():
        raise RuntimeError("LibreOffice no pudo convertir el DOCX a PDF.")
    return pdf


# ----------------------------------------------------------------------------------------------- job

def _all(query_factory: Callable[[], Any], page: int = 1000) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    while True:
        chunk = query_factory().range(len(rows), len(rows) + page - 1).execute().data or []
        rows += chunk
        if len(chunk) < page:
            return rows


def load_context(client: Any, job: dict[str, Any]) -> dict[str, Any]:
    pid, org = job["project_id"], job["org_id"]

    def sel(table: str, cols: str, **eq: Any) -> list[dict[str, Any]]:
        def q() -> Any:
            b = client.table(table).select(cols)
            for k, v in eq.items():
                b = b.eq(k, v)
            return b
        return _all(q)

    proj = sel("projects", "*, clients(name, cuit, address)", id=pid)[0]
    client_row = proj.pop("clients", None) or {}
    codes = {c["code"]: c["meaning"] for c in sel("catalog_codes", "code, meaning", org_id=org)}
    fichas = {l["id"]: l["ficha_no"] for l in sel("survey_lines", "id, ficha_no", project_id=pid)}
    tpl_rows = sel("catalog_text_blocks", "template", org_id=org, scope="interferencia", key="interferencia")
    inter, sin = build_interferencias(
        sel("waypoints_view", "id, line_id, number, code, description, views, lat, lon, elevation_m", project_id=pid),
        codes, fichas, tpl_rows[0]["template"] if tpl_rows else None)
    cats = sel("catalog_photo_categories", "key, label, sort_order", org_id=org)

    # --- contenido de las secciones (catálogos + lo que el profesional eligió o ajustó en este proyecto)
    from app.core.text_template import project_vars, xml_safe_counted

    blocks = sel("catalog_text_blocks", "key, scope, title, template, sort_order", org_id=org)
    penv = {r["item_id"]: r for r in sel("project_environment", "item_id, included, body_override", project_id=pid)}
    environment = []
    for it in (sel("catalog_environment", "id, zone_key, section, label, body, sort_order", org_id=org, zone_key=proj["zone_key"])
               if proj.get("zone_key") else []):
        row = penv.get(it["id"])
        if row is not None and not row["included"]:
            continue
        environment.append({"section": it["section"], "label": it["label"], "sort_order": it["sort_order"],
                            "body": (row or {}).get("body_override") or it["body"]})
    environment.sort(key=lambda x: x["sort_order"] or 0)
    measures = sel("catalog_measures", "id, general, stage, action, resource, timing, responsible, follow_up, body, sort_order", org_id=org)
    chosen = {r["measure_id"]: r for r in sel("project_measures", "measure_id, selected, responsible, timing", project_id=pid) if r["selected"]}
    pga = [{"stage": m["stage"], "action": m["action"], "measure": m["body"], "resource": m["resource"], "follow_up": m["follow_up"],
            "timing": chosen[m["id"]].get("timing") or m["timing"], "responsible": chosen[m["id"]].get("responsible") or m["responsible"],
            "sort_order": m["sort_order"]}
           for m in sorted(measures, key=lambda x: x["sort_order"] or 0) if not m["general"] and m["id"] in chosen]
    wells = []
    for w in sel("wells_geojson", "name, geojson", project_id=pid):
        g = w.get("geojson") or {}
        if g.get("type") == "Point" and g.get("coordinates"):
            wells.append({"name": w["name"], "lon": g["coordinates"][0], "lat": g["coordinates"][1]})
    context, removed_count = xml_safe_counted({
        "vars": project_vars(proj, client_row.get("name")),
        "sections": [b for b in blocks if b["scope"] == "seccion"],
        "declarations": [b for b in blocks if b["scope"] == "declaracion"],
        "section_overrides": {r["key"]: r["body"] for r in sel("project_section_texts", "key, body", project_id=pid)},
        "decl_overrides": {r["factor_id"]: r["body_override"] for r in sel("project_declarations", "factor_id, body_override", project_id=pid)},
        "factors": sel("catalog_impact_factors", "id, code, name, medio, uip, component, sort_order", org_id=org),
        "actions": sel("catalog_impact_actions", "id, code, name, stage, sort_order", org_id=org),
        "impacts": sel("project_impacts", "action_id, factor_id, sign, importance, category", project_id=pid),
        "environment": environment,
        "pga_general": [m["body"] for m in sorted(measures, key=lambda x: x["sort_order"] or 0) if m["general"]],
        "pga": pga,
        "wells": wells,
        "figures": {r["kind"]: r["file_path"] for r in sorted(
            sel("figure_builds", "kind, file_path, created_at, status", project_id=pid, status="listo"), key=lambda x: x["created_at"])
            if r.get("file_path")},
        "final": (job.get("params") or {}).get("final") is True,  # estricto: "false" (texto) no es True
        "_org": job["org_id"],
        "project": proj, "client": client_row,
        "works": sel("works_compare", "name, kind, declared_length_m, declared_area_m2, geom_length_m, geom_area_m2, sort_order", project_id=pid),
        "interferencias": inter, "interferencias_sin_posicion": sin,
        "photos": sel("photos", "id, category, path_original, caption, taken_at", project_id=pid),
        "photo_categories": sorted(cats, key=lambda c: c.get("sort_order") or 0),
        "layers": [{"name": l["base_name"], "kind": str(l["format"]).upper(), "n": l.get("n_features"),
                    "files": [f["path"] for f in (l.get("files") or []) if f.get("path")]}
                   for l in sel("layer_imports", "base_name, format, n_features, status, files", project_id=pid) if l["status"] in ("listo", "incompleto")],
        "gps": [{"name": Path(g["file_path"]).name, "kind": str(g["file_kind"]).upper(), "n": g.get("n_points"), "files": [g["file_path"]]}
                for g in sel("gps_imports", "file_path, file_kind, n_points, status", project_id=pid) if g["status"] == "listo"],
    })
    context["_xml_removed"] = removed_count
    return context


def _own_path(ctx: dict[str, Any], path: str) -> str:
    """El worker usa service role (salta RLS de Storage): solo descarga objetos de la organización del trabajo."""
    if not path.startswith(f"{ctx['_org']}/") or ".." in path.split("/"):
        raise ValueError(f"ruta fuera de la organización: {path}")
    return path


def _load_photos(client: Any, ctx: dict[str, Any], params: dict[str, int], log: list[str]) -> list[dict[str, Any]]:
    order = {c["key"]: i for i, c in enumerate(ctx["photo_categories"])}
    label = {c["key"]: c["label"] for c in ctx["photo_categories"]}
    out: list[dict[str, Any]] = []
    for ph in sorted(ctx["photos"], key=lambda x: (order.get(x["category"], 999), x.get("taken_at") or "")):
        if not ph.get("path_original"):  # subida sin terminar: se informa (y la versión final se rechaza)
            message = f"Foto omitida (sin archivo subido, categoría {ph.get('category')})"
            log.append(message)
            ctx.setdefault("_omitted", []).append(message)
            continue
        try:
            raw = client.storage.from_(BUCKET).download(_own_path(ctx, ph["path_original"]))
            out.append({"category": ph["category"], "label": label.get(ph["category"], ph["category"]),
                        "caption": ph.get("caption"), "jpeg": prepare_photo(raw, params["photo_max_px"], params["jpeg_quality"])})
        except Exception as e:  # en borrador se informa; la versión final falla después de reunir todas las omisiones
            message = f"Foto omitida ({ph['path_original']}): {e}"
            log.append(message)
            ctx.setdefault("_omitted", []).append(message)
    return out


def interferencias_kmz(rows: list[dict[str, Any]]) -> bytes:
    """KMZ con las interferencias (una marca por fila de la tabla). Solo puntos con posición."""
    from xml.sax.saxutils import escape

    marks = []
    for r in rows:
        if r.get("_lat") is None or r.get("_lon") is None:
            continue
        ele = r.get("_ele") or 0
        marks.append(f"<Placemark><name>{escape(r['figura'])}</name><description>{escape(r['descripcion'])}</description>"
                     f"<Point><coordinates>{r['_lon']},{r['_lat']},{ele}</coordinates></Point></Placemark>")
    kml = ('<?xml version="1.0" encoding="UTF-8"?><kml xmlns="http://www.opengis.net/kml/2.2"><Document><name>Interferencias</name>'
           + "".join(marks) + "</Document></kml>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("doc.kml", kml)
    return buf.getvalue()


def make_package(client: Any, ctx: dict[str, Any], docx: bytes, pdf: bytes | None, log: list[str]) -> bytes:
    """Paquete final para enviar por fuera de la app: informe aprobado + KMZ de interferencias + anexos georreferenciados."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("Informe.docx", docx)
        if pdf:
            z.writestr("Informe.pdf", pdf)
        z.writestr("Interferencias.kmz", interferencias_kmz(ctx["interferencias"]))
        for f in ctx["layers"] + ctx["gps"]:
            if not f.get("files"):
                message = f"Paquete: {f.get('name')} figura en el informe pero no tiene archivos para incluir"
                log.append(message)
                ctx.setdefault("_omitted", []).append(message)
            for path in f.get("files") or []:
                try:
                    z.writestr(f"Anexos georreferenciados/{Path(path).name}", client.storage.from_(BUCKET).download(_own_path(ctx, path)))
                except Exception as e:  # se reúnen todos los anexos faltantes antes de rechazar el paquete final
                    message = f"Paquete: no se pudo incluir {Path(path).name} ({e})"
                    log.append(message)
                    ctx.setdefault("_omitted", []).append(message)
    if ctx.get("_omitted"):
        raise ValueError("La versión FINAL no puede generarse porque se omitieron elementos:\n- "
                         + "\n- ".join(ctx["_omitted"]))
    return buf.getvalue()


def run_docs_job(client: Any, job: dict[str, Any], run: Runner = subprocess.run) -> None:
    """Procesa un document_build ya reclamado. Nunca lanza: deja estado y log en la fila."""
    log: list[str] = []
    upd: dict[str, Any]
    try:
        params = clean_params(job.get("params"))
        ctx = load_context(client, job)
        if ctx.get("_xml_removed"):
            log.append(f"Se quitaron {ctx['_xml_removed']} caracteres de control de los textos "
                       "(suelen venir de pegar desde Word).")
        ctx["_omitted"] = []
        photos = _load_photos(client, ctx, params, log)
        ctx["figure_images"] = {}
        for kind, path in (ctx.get("figures") or {}).items():
            try:
                ctx["figure_images"][kind] = client.storage.from_(BUCKET).download(_own_path(ctx, path))
            except Exception as e:  # en borrador se informa; la versión final se rechaza al armar el paquete
                message = f"Figura omitida ({kind}): {e}"
                log.append(message)
                ctx["_omitted"].append(message)
        log.append(f"{len(photos)} foto(s) a {params['photo_max_px']} px, calidad {params['jpeg_quality']}.")
        if ctx.get("final") and ctx["_omitted"]:  # falla antes de armar y subir nada
            raise ValueError("La versión FINAL no puede generarse porque se omitieron elementos:\n- "
                             + "\n- ".join(ctx["_omitted"]))

        if job.get("template_id"):
            trow = client.table("document_templates").select("file_path").eq("id", job["template_id"]).limit(1).execute().data
            if not trow:
                raise ValueError("La plantilla elegida ya no existe.")
            docx_bytes = render_with_template(client.storage.from_(BUCKET).download(_own_path(ctx, trow[0]["file_path"])), ctx, photos)
        else:
            docx_bytes = build_docx(ctx, photos)

        base = f"{job['org_id']}/{job['project_id']}/docs/{job['id']}"
        pdf_path: str | None = None
        pdf_bytes: bytes | None = None
        with tempfile.TemporaryDirectory() as tmp:
            docx = Path(tmp) / "informe.docx"
            docx.write_bytes(docx_bytes)
            try:
                pdf = docx_to_pdf(docx, run)
                pdf_bytes = pdf.read_bytes()
                client.storage.from_(BUCKET).upload(f"{base}/informe.pdf", pdf_bytes,
                                                    {"content-type": "application/pdf", "upsert": "true"})
                pdf_path = f"{base}/informe.pdf"
            except RuntimeError as e:
                log.append(str(e))
        client.storage.from_(BUCKET).upload(
            f"{base}/informe.docx", docx_bytes,
            {"content-type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "upsert": "true"})
        upd = {"status": "listo", "docx_path": f"{base}/informe.docx", "pdf_path": pdf_path}
        log.extend(ctx.get("_warn") or [])
        log.append("Versión FINAL (aprobada): sin marca de borrador." if ctx.get("final") else "Versión BORRADOR: sin aprobar.")
        if ctx.get("final"):
            client.storage.from_(BUCKET).upload(f"{base}/paquete_final.zip", make_package(client, ctx, docx_bytes, pdf_bytes, log),
                                                {"content-type": "application/zip", "upsert": "true"})
            upd["package_path"] = f"{base}/paquete_final.zip"
    except (ValueError, RuntimeError) as e:
        log.append(str(e))
        upd = {"status": "error"}
    except Exception as e:
        log.append(f"Error inesperado al generar el informe: {e}")
        upd = {"status": "error"}
    upd["log"] = "\n".join(log)
    upd["finished_at"] = dt.datetime.now(dt.timezone.utc).isoformat()
    client.table("document_builds").update(upd).eq("id", job["id"]).execute()
