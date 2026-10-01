import datetime as dt
import io
from typing import Any

import pytest
from docx import Document

from app.core.text_template import fill_vars, paragraphs, project_vars, unresolved
from app.jobs.docs import build_docx

HOY = dt.date(2026, 9, 30)


def text_of(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    for s in d.sections:
        parts += [p.text for p in s.header.paragraphs]
    return "\n".join(parts)


def base() -> dict[str, Any]:
    return {
        "project": {"name": "Perforación de 6 pozos en PAD58", "code": "2947-26", "short_name": "PAD 58", "doc_type": "IA",
                    "field_area": "Bajada del Palo Oeste", "province": "Neuquén", "zone_key": "bajada_del_palo_oeste",
                    "applicant": {"razon_social": "Operadora SA"}, "consultant": {"razon_social": "Consultora SRL", "responsable": "Ing. Pérez"}},
        "client": {"name": "Operadora SA"},
        "works": [], "interferencias": [], "interferencias_sin_posicion": 0, "photos": [], "photo_categories": [], "layers": [], "gps": [],
    }


def completo() -> dict[str, Any]:
    c = base()
    c.update({
        "sections": [
            {"key": "sec_resumen", "title": "RESUMEN EJECUTIVO", "template": "El proyecto {pad} perfora 6 pozos.\n\nSegundo párrafo.", "sort_order": 1},
            {"key": "sec_loc", "title": "UBICACIÓN Y DESCRIPCIÓN GENERAL DEL PROYECTO / Localización Física", "template": "Se ubicará en {yacimiento}.", "sort_order": 2},
            {"key": "sec_ref", "title": "REFERENCIAS / Marco Legal", "template": "Ley 25.675.", "sort_order": 3},
        ],
        "section_overrides": {},
        "factors": [
            {"id": "fa", "code": "F01", "name": "Calidad perceptible del aire", "medio": "fisico", "uip": 50, "component": "AIRE", "sort_order": 1},
            {"id": "fb", "code": "F13", "name": "Actividades económicas relacionadas", "medio": "socioeconomico", "uip": 50, "component": "ACT", "sort_order": 2},
        ],
        "actions": [
            {"id": "a1", "code": "A01", "name": "Traslado de equipos y materiales", "stage": "construccion", "sort_order": 1},
            {"id": "a2", "code": "A04", "name": "Operaciones de perforación", "stage": "perforacion", "sort_order": 2},
        ],
        "impacts": [
            {"action_id": "a1", "factor_id": "fa", "sign": -1, "importance": -21, "category": "Bajo"},
            {"action_id": "a2", "factor_id": "fa", "sign": -1, "importance": -28, "category": "Moderado"},
            {"action_id": "a2", "factor_id": "fb", "sign": 1, "importance": 30, "category": "Positivo"},
        ],
        "declarations": [
            {"key": "decl_F01", "title": "Aire – Afectación sobre la calidad perceptible del aire", "template": "La calidad del aire en {pad} podría alterarse."},
            {"key": "decl_F13", "title": "Actividades económicas", "template": "Efecto positivo en {yacimiento}."},
        ],
        "decl_overrides": {},
        "environment": [
            {"section": "geologia", "label": "Geología", "body": "Formación Vaca Muerta en {yacimiento}.", "sort_order": 1},
            {"section": "clima", "label": "Vientos", "body": "Predominan vientos del oeste.", "sort_order": 2},
        ],
        "pga_general": ["Se prohíbe la caza."],
        "pga": [{"stage": "construccion", "action": "Transporte", "measure": "Circular por accesos permitidos en {pad}.",
                 "resource": "Calidad del aire", "timing": "Durante la obra", "responsible": "Supervisores", "follow_up": "Registros", "sort_order": 1}],
        "wells": [{"name": "VIS.Nq.BPO-2581(h)", "lat": -38.13, "lon": -68.57}],
        "final": False,
    })
    return c


# --- variables
def test_fill_vars_paridad_con_la_web() -> None:
    assert fill_vars("Las tareas de {pad} son cortas.\n\nEn {pad}, hay ductos.", {"pad": "PAD 58"}) == \
        "Las tareas de PAD 58 son cortas.\n\nEn PAD 58, hay ductos."
    assert fill_vars("Área {yacimiento} y {rara}", {"yacimiento": ""}) == "Área  y {rara}"
    assert fill_vars("{n} pozos, {x}", {"n": 6, "x": None}) == "6 pozos, "
    assert unresolved("{pad} {foo} {foo} {bar}", {"pad": "x"}) == ["foo", "bar"]
    assert paragraphs("uno\n\n\ndos\n \ntres") == ["uno", "dos", "tres"]


def test_project_vars_cae_al_nombre_si_falta_el_corto() -> None:
    p = base()["project"]
    assert project_vars(p)["pad"] == "PAD 58"
    assert project_vars({**p, "short_name": "  "})["pad"] == p["name"]


# --- informe completo
def test_informe_completo_incluye_todos_los_capitulos_y_reemplaza_variables() -> None:
    t = text_of(build_docx(completo(), [], HOY))
    for esperado in (
        "2. Resumen ejecutivo", "El proyecto PAD 58 perfora 6 pozos.", "Segundo párrafo.",
        "3. Ubicación y descripción general del proyecto", "Localización Física", "Se ubicará en Bajada del Palo Oeste.",
        "Coordenadas de los pozos", "VIS.Nq.BPO-2581(h)",
        "4. Descripción general del ambiente", "Formación Vaca Muerta en Bajada del Palo Oeste.",
        "5. Identificación de impactos y efectos ambientales", "Valoración de los impactos por factor", "−28 Moderado",
        "6. Declaración de impacto ambiental", "La calidad del aire en PAD 58 podría alterarse.", "Efecto positivo en Bajada del Palo Oeste.",
        "7. Plan de gestión ambiental", "Se prohíbe la caza.", "Circular por accesos permitidos en PAD 58.",
        "8. Referencias", "Ley 25.675.", "9. Anexos", "Análisis matricial", "A01", "Traslado de equipos y materiales",
    ):
        assert esperado in t, esperado
    assert "{pad}" not in t and "{yacimiento}" not in t


def test_el_ajuste_del_proyecto_pisa_al_catalogo() -> None:
    c = completo()
    c["section_overrides"] = {"sec_resumen": "Resumen propio de {pad}."}
    c["decl_overrides"] = {"fa": "Declaración ajustada para el aire."}
    t = text_of(build_docx(c, [], HOY))
    assert "Resumen propio de PAD 58." in t and "El proyecto PAD 58 perfora" not in t
    assert "Declaración ajustada para el aire." in t and "podría alterarse" not in t


def test_borrador_lleva_encabezado_y_la_final_no() -> None:
    assert "BORRADOR — versión sin aprobar" in text_of(build_docx(completo(), [], HOY))
    final = completo()
    final["final"] = True
    assert "BORRADOR" not in text_of(build_docx(final, [], HOY))


def test_la_matriz_sale_en_hoja_horizontal_y_el_resto_en_vertical() -> None:
    d = Document(io.BytesIO(build_docx(completo(), [], HOY)))
    orient = [s.orientation for s in d.sections]
    assert 1 in orient                  # WD_ORIENT.LANDSCAPE == 1: hay una hoja horizontal (la matriz)
    assert d.sections[0].orientation == 0 and d.sections[-1].orientation == 0   # empieza y termina en vertical


def test_secciones_sin_datos_se_marcan_incompletas_sin_inventar() -> None:
    c = base()
    c.update({"sections": [], "factors": [], "actions": [], "impacts": [], "declarations": [], "environment": [],
              "pga_general": [], "pga": [], "wells": [], "section_overrides": {}, "decl_overrides": {}})
    c["project"]["zone_key"] = None
    t = text_of(build_docx(c, [], HOY))
    for esperado in ("no hay texto de resumen ejecutivo", "no se eligió la zona del proyecto", "todavía no tiene celdas cargadas",
                     "no hay declaraciones en el catálogo", "no se eligieron medidas del PGA", "no hay referencias ni marco legal"):
        assert esperado in t, esperado


def test_variable_desconocida_queda_visible_y_se_avisa() -> None:
    c = completo()
    c["sections"][0]["template"] = "Texto con {variable_rara} adentro."
    c["_warn"] = []
    t = text_of(build_docx(c, [], HOY))
    assert "{variable_rara}" in t
    assert any("variable_rara" in w for w in c["_warn"])


# --- paquete final y job
import zipfile  # noqa: E402
from types import SimpleNamespace  # noqa: E402

from app.jobs import docs  # noqa: E402
from app.jobs.docs import interferencias_kmz, make_package, run_docs_job  # noqa: E402


def test_kmz_de_interferencias() -> None:
    rows = [{"figura": "Cruce & cauce", "descripcion": "con <ductos>", "_lat": -38.13, "_lon": -68.57, "_ele": 138.0},
            {"figura": "sin posición", "descripcion": "x", "_lat": None, "_lon": None, "_ele": None}]
    z = zipfile.ZipFile(io.BytesIO(interferencias_kmz(rows)))
    kml = z.read("doc.kml").decode()
    assert kml.count("<Placemark>") == 1                        # solo con posición
    assert "Cruce &amp; cauce" in kml and "con &lt;ductos&gt;" in kml   # XML bien escapado
    assert "<coordinates>-68.57,-38.13,138.0</coordinates>" in kml      # lon,lat,cota


class Store:
    def __init__(self, files: dict[str, bytes]) -> None:
        self.files, self.uploads = files, {}
        self.storage = self
    def from_(self, _b: str): return self
    def download(self, p: str) -> bytes:
        if p not in self.files:
            raise FileNotFoundError(p)
        return self.files[p]
    def upload(self, p: str, data: bytes, _o: Any): self.uploads[p] = data


def test_paquete_final_reune_informe_kmz_y_anexos() -> None:
    c = completo()
    c["interferencias"] = [{"figura": "Cruce", "descripcion": "d", "_lat": -38.1, "_lon": -68.5, "_ele": None}]
    c["layers"] = [{"name": "Caminos", "kind": "SHP", "n": 5, "files": ["o/p/layers/i1/Caminos.shp", "o/p/layers/i1/Caminos.dbf"]}]
    c["gps"] = [{"name": "PAD58.gdb", "kind": "GDB", "n": 30, "files": ["o/p/gps/g1/PAD58.gdb"]}]
    log: list[str] = []
    c["_org"] = "o"
    store = Store({"o/p/layers/i1/Caminos.shp": b"shp", "o/p/layers/i1/Caminos.dbf": b"dbf",
                   "o/p/gps/g1/PAD58.gdb": b"gdb"})

    from app.jobs.docs import _own_path
    assert _own_path(c, "o/p/x.jpg") == "o/p/x.jpg"
    for malo in ("otra/p/x.jpg", "o/../otra/x.jpg"):   # ruta de otra organización o con ..
        try:
            _own_path(c, malo)
        except ValueError:
            continue
        raise AssertionError(malo)
    z = zipfile.ZipFile(io.BytesIO(make_package(store, c, b"DOCX", b"%PDF", log)))
    assert sorted(z.namelist()) == ["Anexos georreferenciados/Caminos.dbf", "Anexos georreferenciados/Caminos.shp",
                                    "Anexos georreferenciados/PAD58.gdb",
                                    "Informe.docx", "Informe.pdf", "Interferencias.kmz"]
    assert log == []


def test_paquete_final_falla_listando_todas_las_omisiones() -> None:
    c = completo()
    c.update({
        "final": True,
        "_org": "o",
        "_omitted": ["Foto omitida (o/p/photos/roto.jpg): falta"],
        "interferencias": [],
        "layers": [{"files": ["o/p/layers/Caminos.dbf", "o/p/layers/Caminos.shx"]}],
        "gps": [],
    })
    log: list[str] = []

    with pytest.raises(ValueError) as error:
        make_package(Store({}), c, b"DOCX", None, log)

    message = str(error.value)
    for item in ("o/p/photos/roto.jpg", "Caminos.dbf", "Caminos.shx"):
        assert item in message


def test_solo_la_version_final_genera_paquete(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(docs, "find_soffice", lambda: None)
    llamadas: list[bool] = []
    monkeypatch.setattr(docs, "load_context", lambda client, job: {**completo(), "interferencias": [], "final": bool(job["params"].get("final")),
                                                                    "works": [], "interferencias_sin_posicion": 0,
                                                                    "photos": [], "photo_categories": [], "layers": [], "gps": []})

    class C(Store):
        def __init__(self) -> None:
            super().__init__({})
            self.updates: list[dict[str, Any]] = []
            self._u: Any = None
        def table(self, _t: str): return self
        def update(self, u: dict[str, Any]): self._u = u; return self
        def eq(self, *_a): return self
        def execute(self):
            self.updates.append(self._u)
            return SimpleNamespace(data=[])

    borrador, final = C(), C()
    run_docs_job(borrador, {"id": "b1", "org_id": "o", "project_id": "p", "template_id": None, "params": {}})
    run_docs_job(final, {"id": "b2", "org_id": "o", "project_id": "p", "template_id": None, "params": {"final": True}})
    assert "package_path" not in borrador.updates[-1] and "BORRADOR" in borrador.updates[-1]["log"]
    assert final.updates[-1]["package_path"] == "o/p/docs/b2/paquete_final.zip" and "FINAL" in final.updates[-1]["log"]
    assert "o/p/docs/b2/paquete_final.zip" in final.uploads and "o/p/docs/b1/paquete_final.zip" not in borrador.uploads
    llamadas.append(True)


# --- plantillas de cliente
def test_datos_para_plantillas_traen_todo_el_contenido_ya_resuelto() -> None:
    from app.core.report_sections import template_data
    from app.core.text_template import project_vars

    c = completo()
    c["vars"] = project_vars(c["project"], "Operadora SA")
    d = template_data(c)
    assert d["borrador"] is True
    assert d["secciones"]["resumen"][0]["texto"].startswith("El proyecto PAD 58 perfora")
    assert d["ambiente"][0]["texto"] == "Formación Vaca Muerta en Bajada del Palo Oeste."
    assert d["declaraciones"][0]["texto"] == "La calidad del aire en PAD 58 podría alterarse."
    assert d["pga"]["generales"] == ["Se prohíbe la caza."]
    assert d["pga"]["particulares"][0]["medida"] == "Circular por accesos permitidos en PAD 58."
    assert d["matriz"]["acciones"][0]["codigo"] == "A01"
    fila = next(f for f in d["matriz"]["filas"] if f["factor"] == "Calidad perceptible del aire")
    assert fila["valores"] == ["−21", "−28"]
    assert next(f for f in d["factores"] if f["codigo"] == "F01")["categoria_peor"] == "Moderado"
    assert d["pozos"][0]["nombre"] == "VIS.Nq.BPO-2581(h)"
    c["final"] = True
    assert template_data(c)["borrador"] is False


def test_plantilla_de_cliente_se_completa_con_los_marcadores_nuevos() -> None:
    from app.jobs.docs import render_with_template

    tpl = Document()
    tpl.add_paragraph("{{ titulo }} — {{ proyecto.name }} ({{ variables.pad }})")
    tpl.add_paragraph("{%p if borrador %}")
    tpl.add_paragraph("BORRADOR")
    tpl.add_paragraph("{%p endif %}")
    tpl.add_paragraph("{%p for d in declaraciones %}")
    tpl.add_paragraph("{{ d.titulo }}: {{ d.texto }}")
    tpl.add_paragraph("{%p endfor %}")
    tpl.add_paragraph("{%p for m in pga.particulares %}")
    tpl.add_paragraph("{{ m.etapa }} / {{ m.medida }}")
    tpl.add_paragraph("{%p endfor %}")
    buf = io.BytesIO()
    tpl.save(buf)

    out = text_of(render_with_template(buf.getvalue(), completo(), []))
    assert "INFORME AMBIENTAL — Perforación de 6 pozos en PAD58 (PAD 58)" in out
    assert "BORRADOR" in out
    assert "Aire – Afectación sobre la calidad perceptible del aire: La calidad del aire en PAD 58 podría alterarse." in out
    assert "Construcción / Circular por accesos permitidos en PAD 58." in out
    final = completo()
    final["final"] = True
    assert "BORRADOR" not in text_of(render_with_template(buf.getvalue(), final, []))


# --- figuras
def _png() -> bytes:
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 120, 40)).save(b, "PNG")
    return b.getvalue()


def test_las_figuras_generadas_se_incrustan_con_epigrafe_numerado() -> None:
    c = completo()
    c["figure_images"] = {"ubicacion": _png(), "implantacion": _png(), "interferencias": _png()}
    d = Document(io.BytesIO(build_docx(c, [], HOY)))
    t = text_of(build_docx(c, [], HOY))
    assert len(d.inline_shapes) == 3
    for esperado in ("Figura 1. Ubicación general del proyecto.", "Figura 2. Implantación de las obras.", "Figura 3. Interferencias relevadas."):
        assert esperado in t, esperado


def test_sin_figuras_no_agrega_nada_ni_rompe() -> None:
    d = Document(io.BytesIO(build_docx(completo(), [], HOY)))
    assert len(d.inline_shapes) == 0
    assert "Figura 1" not in text_of(build_docx(completo(), [], HOY))


def test_texto_con_caracteres_de_control_no_rompe_el_informe() -> None:
    """Pegado desde Word puede traer \x0b/\x00: python-docx lanzaría ValueError y se perdería todo el informe."""
    from app.core.text_template import xml_safe

    c = xml_safe(completo())
    c["sections"][0]["template"] = xml_safe("Texto con\x0b salto\x00 vertical de {pad}.")
    assert "\x0b" not in c["sections"][0]["template"]
    t = text_of(build_docx(c, [], HOY))
    assert "Texto con salto vertical de PAD 58." in t
    assert xml_safe({"a": ["x\x01y", {"b": "z\x1f"}]}) == {"a": ["xy", {"b": "z"}]}
    assert xml_safe("tab\tsalto\nok") == "tab\tsalto\nok"       # tab y salto de línea sí son válidos


def test_xml_safe_counted_devuelve_texto_limpio_y_cantidad() -> None:
    from app.core.text_template import xml_safe_counted

    limpio, quitados = xml_safe_counted({"a": ["x\x01y\x02", {"b": "z\x1f"}], "valido": "tab\ty\nsalto"})

    assert limpio == {"a": ["xy", {"b": "z"}], "valido": "tab\ty\nsalto"}
    assert quitados == 3
