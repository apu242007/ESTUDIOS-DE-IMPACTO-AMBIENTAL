import datetime as dt
import io
import os
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest
from docx import Document
from PIL import Image

from app.core.geo import format_dms, to_gauss_kruger
from app.jobs import docs
from app.jobs.docs import (
    build_docx, build_interferencias, clean_params, expand_codes, docx_to_pdf, prepare_photo, render_template, run_docs_job,
)

# Punto de A.7/A.8 (mismo que la web): lat -38°7'47.78", lon -68°34'8.97"
LAT = -(38 + 7 / 60 + 47.78 / 3600)
LON = -(68 + 34 / 60 + 8.97 / 3600)


def jpeg(w: int = 3000, h: int = 2000) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (120, 160, 90)).save(buf, "JPEG")
    return buf.getvalue()


def ctx(**over: Any) -> dict[str, Any]:
    base: dict[str, Any] = {
        "project": {"name": "PAD 58", "code": "2947-26", "doc_type": "IA", "field_area": "Loma Campana", "province": "Neuquén",
                    "applicant": {"razon_social": "Operadora SA", "cuit": "30-1-2", "domicilio": "Calle 1"},
                    "consultant": {"razon_social": "Consultora SRL", "responsable": "Ing. Pérez", "matricula": "123"}},
        "client": {"name": "Operadora SA"},
        "works": [{"name": "Camino troncal", "kind": "camino", "declared_length_m": 2270, "declared_area_m2": None,
                   "geom_length_m": 2269.2, "geom_area_m2": None}],
        "interferencias": [{"figura": "Cruce", "lat": "38° 7'47.78\"S", "lon": "68°34'8.97\"O", "x": 5779957, "y": 2537775,
                            "cota": 138, "descripcion": "Cruce: con ductos (4)"}],
        "interferencias_sin_posicion": 0, "photos": [], "photo_categories": [], "layers": [], "gps": [],
    }
    base.update(over)
    return base


def text_of(data: bytes) -> str:
    d = Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts += [c.text for c in row.cells]
    return "\n".join(parts)


# --- coordenadas: mismos resultados y formato que la web
def test_gauss_kruger_paridad_con_la_web() -> None:
    x, y = to_gauss_kruger(LAT, LON)
    assert x == pytest.approx(5779957.1, abs=0.05) and y == pytest.approx(2537775.1, abs=0.05)
    assert x > y  # X es el norte
    x0, y0 = to_gauss_kruger(0, -69)
    assert x0 == pytest.approx(10001965.729, abs=0.01) and y0 == pytest.approx(2500000, abs=0.001)


def test_dms_paridad_con_la_web() -> None:
    assert format_dms(LAT, LON) == ('38° 7\'47.78"S', '68°34\'8.97"O')
    assert format_dms(10.5, 20.25) == ('10°30\'0.00"N', '20°15\'0.00"E')
    assert format_dms(0.99999999, 0)[0] == '1° 0\'0.00"N'


# --- interferencias
def test_interferencias_ordena_y_descarta_sin_posicion() -> None:
    wps = [
        {"id": "b", "line_id": "l2", "number": 1, "code": "CR", "description": "con ductos", "views": "O-SO", "lat": LAT, "lon": LON, "elevation_m": 138.4},
        {"id": "a", "line_id": "l1", "number": 5, "code": None, "description": "Inicio de línea", "views": None, "lat": LAT, "lon": LON, "elevation_m": None},
        {"id": "c", "line_id": "l1", "number": 6, "code": "CR", "description": None, "views": None, "lat": None, "lon": None, "elevation_m": None},
    ]
    rows, sin = build_interferencias(wps, {"CR": "Cruce"}, {"l1": 1, "l2": 2})
    assert sin == 1 and [r["figura"] for r in rows] == ["Punto de interés", "Cruce"]
    assert rows[1]["descripcion"] == "Cruce: con ductos" and rows[1]["cota"] == 138  # las vistas van al anexo de fotos
    assert abs(rows[1]["x"] - 5779960) < 5 and abs(rows[1]["y"] - 2537773) < 5


def test_interferencias_formato_cliente_filtra_quiebres_y_expande_siglas() -> None:
    codes = {"CR": "Cruce", "CP": "Camino principal", "CaC": "Caño camisa", "Q": "Quiebre", "O": "Oleoducto"}
    base = {"line_id": "l1", "views": "O-SO", "lat": LAT, "lon": LON, "elevation_m": None}
    wps = [
        {**base, "id": "1", "number": 9, "code": "CR", "description": "CR con CP - CaC"},
        {**base, "id": "2", "number": 10, "code": "Q", "description": "Q al O"},             # quiebre: fuera
        {**base, "id": "3", "number": 4, "code": None, "description": "Inicio en PAD 60 BPO"},  # punto de interés
        {**base, "id": "4", "number": 50, "code": None, "description": "VNO"},               # nota de campo: fuera
        {**base, "id": "5", "number": 51, "code": None, "description": None},                # renglón vacío: fuera
    ]
    rows, _ = build_interferencias(wps, codes, {"l1": (1, "Acueductos flexibles")})
    assert [r["_num"] for r in rows] == [4, 9]
    assert rows[1]["descripcion"] == "Cruce con camino principal - caño camisa"
    assert rows[0]["descripcion"] == "Inicio en PAD 60 BPO" and rows[0]["_traza"] == "Acueductos flexibles"
    assert expand_codes("Q al O", codes) == "Q al O"  # siglas de una letra (rumbos) no se tocan


def test_plantilla_de_texto() -> None:
    assert render_template("{figura} {observaciones}.", {"figura": "Cruce"}) == "Cruce."


# --- parámetros y fotos
def test_params_acotados() -> None:
    assert clean_params(None) == {"photo_max_px": 1200, "jpeg_quality": 75}
    assert clean_params({"photo_max_px": 99999, "jpeg_quality": 1}) == {"photo_max_px": 3000, "jpeg_quality": 40}
    assert clean_params({"photo_max_px": "x"})["photo_max_px"] == 1200


def test_prepare_photo_reduce_sin_agrandar() -> None:
    out = Image.open(io.BytesIO(prepare_photo(jpeg(3000, 2000), 1200, 75)))
    assert max(out.size) == 1200 and out.format == "JPEG"
    small = Image.open(io.BytesIO(prepare_photo(jpeg(400, 300), 1200, 75)))
    assert small.size == (400, 300)  # no agranda
    assert len(prepare_photo(jpeg(3000, 2000), 1200, 75)) < len(jpeg(3000, 2000))


# --- DOCX base
def test_docx_completo_tiene_las_secciones_y_los_datos() -> None:
    photos = [{"category": "locacion", "label": "Locación", "caption": "Vista general", "jpeg": prepare_photo(jpeg(), 800, 70)}]
    data = build_docx(ctx(layers=[{"name": "Caminos", "kind": "SHP", "n": 5}]), photos, dt.date(2026, 9, 30))
    t = text_of(data)
    for esperado in ("INFORME AMBIENTAL", "PAD 58", "Septiembre de 2026", "1. Datos generales", "Operadora SA",
                     "Alcance de obras", "Camino troncal", "2.270,0 m", "Interferencias y puntos de interés",
                     "38° 7'47.78\"S", "5779957", "9. Anexos", "Relevamiento fotográfico", "Foto 1. Vista general", "Caminos"):
        assert esperado in t, esperado
    assert "2.269,2 m" not in t  # lo medido es control interno: no va al informe del cliente
    assert Document(io.BytesIO(data)).inline_shapes  # la foto quedó incrustada


def test_docx_vacio_marca_secciones_incompletas_y_no_inventa() -> None:
    c = ctx(works=[], interferencias=[], project={**ctx()["project"], "applicant": {}, "consultant": {}})
    t = text_of(build_docx(c, [], dt.date(2026, 9, 30)))
    assert t.count("Sección incompleta") >= 5
    assert "faltan razón social del solicitante, consultora, responsable técnico" in t
    assert "no se cargaron las obras" in t and "no hay fotos" in t and "no hay capas ni archivos de GPS" in t


def test_docx_tipos_de_documento() -> None:
    assert "MEMORIA TÉCNICA DESCRIPTIVA" in text_of(build_docx(ctx(project={**ctx()["project"], "doc_type": "MTD"}), []))
    assert "INFORME AMBIENTAL Y MEMORIA" in text_of(build_docx(ctx(project={**ctx()["project"], "doc_type": "IA+MTD"}), []))


# --- PDF
def test_pdf_sin_libreoffice_da_mensaje_claro(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(docs, "find_soffice", lambda: None)
    (tmp_path / "informe.docx").write_bytes(b"x")
    with pytest.raises(RuntimeError, match="LibreOffice no está instalado"):
        docx_to_pdf(tmp_path / "informe.docx")


def test_pdf_arma_el_comando(tmp_path: Path) -> None:
    visto: dict[str, Any] = {}

    def run(cmd: list[str], **kw: Any):
        visto["cmd"] = cmd
        Path(tmp_path / "informe.pdf").write_bytes(b"%PDF")
        return SimpleNamespace(returncode=0)

    (tmp_path / "informe.docx").write_bytes(b"x")
    pdf = docx_to_pdf(tmp_path / "informe.docx", run, exe="soffice")
    assert pdf.name == "informe.pdf"
    assert visto["cmd"][:3] == ["soffice", "--headless", "--convert-to"] and visto["cmd"][3].startswith("pdf:writer_pdf_Export:")


# --- job completo con un cliente falso
class FakeClient:
    def __init__(self, tables: dict[str, list[dict[str, Any]]], files: dict[str, bytes]) -> None:
        self.tables, self.files = tables, files
        self.uploads: dict[str, bytes] = {}
        self.updates: list[dict[str, Any]] = []
        self.storage = self
        self._t = ""
        self._cols = ""
        self._eq: dict[str, Any] = {}
        self._upd: dict[str, Any] | None = None
        self._r = 0

    def from_(self, _b: str): return self
    def download(self, p: str) -> bytes:
        if p not in self.files:
            raise FileNotFoundError(p)
        return self.files[p]
    def upload(self, p: str, data: bytes, _o: dict[str, str]): self.uploads[p] = data

    def table(self, t: str): self._t, self._eq, self._upd, self._r = t, {}, None, 0; return self
    def select(self, cols: str): self._cols = cols; return self
    def eq(self, k: str, v: Any): self._eq[k] = v; return self
    def limit(self, _n: int): return self
    def range(self, a: int, _b: int): self._r = a; return self
    def update(self, u: dict[str, Any]): self._upd = u; return self

    def execute(self):
        r = SimpleNamespace()
        if self._upd is not None:
            self.updates.append(self._upd)
            r.data = []
            return r
        rows = [x for x in self.tables.get(self._t, []) if all(x.get(k) == v for k, v in self._eq.items())]
        r.data = rows if self._r == 0 else []
        return r


def base_tables() -> dict[str, list[dict[str, Any]]]:
    c = ctx()
    return {
        "projects": [{"id": "p", **c["project"], "clients": {"name": "Operadora SA", "cuit": None, "address": None}}],
        "catalog_codes": [{"code": "CR", "meaning": "Cruce", "org_id": "o"}],
        "survey_lines": [{"id": "l1", "ficha_no": 1, "project_id": "p"}],
        "waypoints_view": [{"id": "w1", "line_id": "l1", "number": 4, "code": "CR", "description": "con ductos", "views": "O-SO",
                            "lat": LAT, "lon": LON, "elevation_m": 138.0, "project_id": "p"}],
        "catalog_photo_categories": [{"key": "locacion", "label": "Locación", "sort_order": 1, "org_id": "o"}],
        "works_compare": [{**c["works"][0], "sort_order": 1, "project_id": "p"}],
        "photos": [{"id": "f1", "category": "locacion", "path_original": "o/p/photos/f1.jpg", "caption": "Vista", "taken_at": "2026-03-10", "project_id": "p"},
                   {"id": "f2", "category": "locacion", "path_original": "o/p/photos/roto.jpg", "caption": None, "taken_at": "2026-03-11", "project_id": "p"}],
        "layer_imports": [{"base_name": "Caminos", "format": "shp", "n_features": 5, "status": "listo", "project_id": "p"}],
        "gps_imports": [{"file_path": "o/p/gps/x/PAD58.gdb", "file_kind": "gdb", "n_points": 30, "status": "listo", "project_id": "p"}],
    }


def job() -> dict[str, Any]:
    return {"id": "b1", "org_id": "o", "project_id": "p", "template_id": None, "params": {"photo_max_px": 800}}


def sin_soffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(docs, "find_soffice", lambda: None)


def test_job_completo_sube_docx_y_deja_log(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    c = FakeClient(base_tables(), {"o/p/photos/f1.jpg": jpeg()})
    run_docs_job(c, job())
    up = c.updates[-1]
    assert up["status"] == "listo" and up["docx_path"] == "o/p/docs/b1/informe.docx" and up["pdf_path"] is None
    assert "LibreOffice no está instalado" in up["log"]                 # el PDF no se pudo, se dice
    assert "Foto omitida (o/p/photos/roto.jpg)" in up["log"]            # una foto rota no tira el informe
    assert "1 foto(s) a 800 px" in up["log"] and up["finished_at"]
    t = text_of(c.uploads["o/p/docs/b1/informe.docx"])
    for esperado in ("PAD 58", "Cruce: con ductos", "PAD58.gdb", "Caminos", "Foto 1. Vista"):
        assert esperado in t


def test_job_final_con_foto_faltante_falla_y_lista_el_item(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    c = FakeClient(base_tables(), {"o/p/photos/f1.jpg": jpeg()})
    final = {**job(), "params": {**job()["params"], "final": True}}

    run_docs_job(c, final)

    up = c.updates[-1]
    assert up["status"] == "error"
    assert "La versión FINAL no puede generarse porque se omitieron elementos" in up["log"]
    assert "o/p/photos/roto.jpg" in up["log"]
    assert "Foto omitida" in up["log"]
    assert not any(p.startswith("o/p/docs/") for p in c.uploads)    # falla antes de subir nada


def _foto_sin_archivo() -> dict[str, Any]:
    return {"id": "ph9", "category": "locacion", "path_original": None, "caption": None, "taken_at": None, "project_id": "p"}


def test_job_final_con_foto_sin_archivo_subido_falla(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    t = base_tables()
    t["photos"] = [t["photos"][0], _foto_sin_archivo()]
    c = FakeClient(t, {"o/p/photos/f1.jpg": jpeg(), "o/p/photos/roto.jpg": jpeg()})

    run_docs_job(c, {**job(), "params": {**job()["params"], "final": True}})

    up = c.updates[-1]
    assert up["status"] == "error" and "sin archivo subido" in up["log"]


def test_job_borrador_con_foto_sin_archivo_sigue_listo_y_avisa(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    t = base_tables()
    t["photos"] = [t["photos"][0], _foto_sin_archivo()]
    c = FakeClient(t, {"o/p/photos/f1.jpg": jpeg(), "o/p/photos/roto.jpg": jpeg()})

    run_docs_job(c, job())

    up = c.updates[-1]
    assert up["status"] == "listo" and "sin archivo subido" in up["log"]


def test_final_con_final_texto_no_cuenta_como_final(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    c = FakeClient(base_tables(), {"o/p/photos/f1.jpg": jpeg()})

    run_docs_job(c, {**job(), "params": {**job()["params"], "final": "true"}})   # texto, no booleano: es borrador

    up = c.updates[-1]
    assert up["status"] == "listo" and "package_path" not in up


def test_job_informa_caracteres_de_control_quitados(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    tables = base_tables()
    tables["photos"][0]["caption"] = "Vista\x0b general"
    c = FakeClient(tables, {"o/p/photos/f1.jpg": jpeg()})

    run_docs_job(c, job())

    assert c.updates[-1]["status"] == "listo"
    assert "Se quitaron 1 caracteres de control de los textos" in c.updates[-1]["log"]


def test_job_con_pdf(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(docs, "find_soffice", lambda: "soffice")

    def run(cmd: list[str], **kw: Any):
        Path(cmd[-1]).with_suffix(".pdf").write_bytes(b"%PDF-1.4")
        return SimpleNamespace(returncode=0)

    c = FakeClient(base_tables(), {"o/p/photos/f1.jpg": jpeg()})
    run_docs_job(c, job(), run)
    assert c.updates[-1]["pdf_path"] == "o/p/docs/b1/informe.pdf" and c.uploads["o/p/docs/b1/informe.pdf"] == b"%PDF-1.4"


def test_plantilla_inexistente_da_error_legible() -> None:
    c = FakeClient(base_tables(), {})
    run_docs_job(c, {**job(), "template_id": "t9"})
    up = c.updates[-1]
    assert up["status"] == "error" and "plantilla elegida ya no existe" in up["log"]


def test_proyecto_inexistente_no_lanza() -> None:
    t = base_tables()
    t["projects"] = []
    c = FakeClient(t, {})
    run_docs_job(c, job())
    assert c.updates[-1]["status"] == "error"


def test_job_achica_fotos_si_el_docx_supera_el_limite(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    buf = io.BytesIO()
    Image.frombytes("RGB", (1600, 1200), os.urandom(1600 * 1200 * 3)).save(buf, "JPEG", quality=95)  # ruido: no comprime
    files = {"o/p/photos/f1.jpg": buf.getvalue()}
    c = FakeClient(base_tables(), files)
    run_docs_job(c, job())
    tamano = len(c.uploads["o/p/docs/b1/informe.docx"])
    monkeypatch.setattr(docs, "MAX_UPLOAD", tamano - 1)  # con las fotos a 800 px ya no entra
    c = FakeClient(base_tables(), files)
    run_docs_job(c, job())
    up = c.updates[-1]
    assert up["status"] == "listo" and "fotos reducidas a 640 px" in up["log"]
    assert len(c.uploads["o/p/docs/b1/informe.docx"]) < tamano


def test_job_pone_encabezado_de_la_consultora_y_logo_del_cliente(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    t = base_tables()
    t["organizations"] = [{"id": "o", "header_image_path": "o/branding/h.jpg"}]
    t["projects"][0]["clients"]["logo_path"] = "o/logos/c.jpg"
    c = FakeClient(t, {"o/p/photos/f1.jpg": jpeg(), "o/branding/h.jpg": jpeg(), "o/logos/c.jpg": jpeg()})
    run_docs_job(c, job())
    assert c.updates[-1]["status"] == "listo"
    d = Document(io.BytesIO(c.uploads["o/p/docs/b1/informe.docx"]))
    assert d.sections[0].header._element.xpath(".//*[local-name()='blip']")  # imagen en el encabezado
    assert "BORRADOR" in "".join(p.text for p in d.sections[0].header.paragraphs)
    assert len(d.inline_shapes) >= 2  # logo del cliente en la carátula + la foto


def test_job_sin_logos_avisa_y_sale_igual(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    c = FakeClient(base_tables(), {"o/p/photos/f1.jpg": jpeg()})
    run_docs_job(c, job())
    up = c.updates[-1]
    assert up["status"] == "listo" and "Sin encabezado de la consultora" in up["log"] and "Sin logo del cliente" in up["log"]


def test_fecha_de_la_caratula_sale_del_proyecto() -> None:
    assert docs.report_date({"report_date": "2026-09-15"}) == dt.date(2026, 9, 15)
    assert docs.report_date({"report_date": None}) == dt.date.today()
    c = ctx()
    c["project"] = {**c["project"], "report_date": "2026-09-15", "code": "2947-26"}
    d = Document(io.BytesIO(build_docx(c, [])))
    assert "Septiembre de 2026" in text_of(build_docx(c, []))
    assert "Trabajo Nº 2947-26." in "".join(p.text for p in d.sections[0].footer.paragraphs)


def test_paquete_final_mas_grande_que_el_limite_da_mensaje_claro(monkeypatch: pytest.MonkeyPatch) -> None:
    sin_soffice(monkeypatch)
    monkeypatch.setattr(docs, "make_package", lambda *_a: b"x" * (docs.MAX_UPLOAD + 1))
    c = FakeClient(base_tables(), {"o/p/photos/f1.jpg": jpeg(), "o/p/photos/roto.jpg": jpeg()})

    run_docs_job(c, {**job(), "params": {**job()["params"], "final": True}})

    up = c.updates[-1]
    assert up["status"] == "error"
    assert "paquete final pesa" in up["log"] and "Error inesperado" not in up["log"]
    assert "o/p/docs/b1/paquete_final.zip" not in c.uploads


def test_achicar_fotos_nunca_sube_la_calidad() -> None:
    assert docs.shrink_photos(1000, 75) == (800, 65)
    assert docs.shrink_photos(1000, 55) == (800, 50)
    assert docs.shrink_photos(1000, 40) == (800, 40)   # el usuario eligió 40: no se sube a 50


def test_fichas_sin_numero_no_se_intercalan() -> None:
    base = {"views": None, "lat": LAT, "lon": LON, "elevation_m": None, "code": "CR", "description": None}
    wps = [{**base, "id": f"{l}{n}", "line_id": l, "number": n} for n in (1, 2, 3) for l in ("a", "b")]
    rows, _ = build_interferencias(wps, {"CR": "Cruce"}, {"a": (None, "Ducto"), "b": (None, "Camino")})
    trazas = [r["_traza"] for r in rows]
    assert trazas == ["Ducto"] * 3 + ["Camino"] * 3 or trazas == ["Camino"] * 3 + ["Ducto"] * 3
